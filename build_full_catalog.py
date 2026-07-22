import json
from translations import TRANSLATIONS
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

rows=json.load(open("rows.json"))
for i,r in enumerate(rows):
    r["name_en"]=TRANSLATIONS[r["name_bg"]]
    r["_orig"]=i
# stable sort by brand then original page order
brand_order=["Alpenzu","Cantarelli","Dalla Costa","Deliziosa","Deseo","Foresti","Ibis",
 "Maletti","Olivieri 1882","Terre di Puglia","Villa Grimelli","Villani","Riso Gallo",
 "Iposea","Divella","Polselli","Mancini"]
rows.sort(key=lambda r:(brand_order.index(r["brand"]), r["_orig"]))

wb=Workbook(); ws=wb.active; ws.title="Catalog"
headers=["#","Brand","ART № (SKU)","Product Name (EN)","Product Name (BG)","Weight",
         "Price excl. VAT (EUR)","Price incl. VAT (EUR)","Notes"]
ws.append(headers)
hdr_fill=PatternFill("solid",fgColor="1F4E78"); hdr_font=Font(name="Arial",bold=True,color="FFFFFF",size=11)
thin=Side(style="thin",color="D0D7E2"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
for c in ws[1]:
    c.fill=hdr_fill; c.font=hdr_font; c.border=border
    c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)

brand_colors={}
palette=["FFFFFF","F2F6FB"]
prev_brand=None; band=0
for i,r in enumerate(rows,start=2):
    if r["brand"]!=prev_brand:
        band^=1; prev_brand=r["brand"]
    fill=palette[band]
    vals=[i-1,r["brand"],r["art"],r["name_en"],r["name_bg"],r["weight"],r["p_no"],r["p_with"],r["note"]]
    for col,v in enumerate(vals,start=1):
        cell=ws.cell(i,col,v)
        cell.font=Font(name="Arial",size=10)
        cell.border=border
        cell.fill=PatternFill("solid",fgColor=fill)
        if col in (4,5,9):
            cell.alignment=Alignment(vertical="center",wrap_text=True)
        elif col in (7,8):
            cell.number_format='#,##0.00 "EUR"'
            cell.alignment=Alignment(horizontal="right",vertical="center")
        else:
            cell.alignment=Alignment(horizontal="center",vertical="center")
    if r["brand"]:
        ws.cell(i,2).font=Font(name="Arial",size=10,bold=True)

widths=[5,15,14,44,44,10,20,20,40]
for idx,w in enumerate(widths,start=1):
    ws.column_dimensions[get_column_letter(idx)].width=w
ws.freeze_panes="A2"
ws.auto_filter.ref=f"A1:{get_column_letter(len(headers))}{len(rows)+1}"

# summary sheet
ws2=wb.create_sheet("Summary")
ws2.append(["Brand","Products","Min price incl. VAT (EUR)","Max price incl. VAT (EUR)"])
for c in ws2[1]:
    c.fill=hdr_fill; c.font=hdr_font; c.alignment=Alignment(horizontal="center")
from collections import defaultdict
agg=defaultdict(list)
for r in rows: agg[r["brand"]].append(r["p_with"])
for b in brand_order:
    if b in agg:
        ws2.append([b,len(agg[b]),min(agg[b]),max(agg[b])])
ws2.append(["TOTAL",len(rows),"",""])
for row in ws2.iter_rows(min_row=2):
    row[0].font=Font(name="Arial",size=10,bold=(row[0].value=="TOTAL"))
    for c in row[1:]:
        c.font=Font(name="Arial",size=10); c.alignment=Alignment(horizontal="center")
        if c.column in (3,4) and isinstance(c.value,(int,float)): c.number_format='#,##0.00'
for idx,w in enumerate([16,12,26,26],start=1):
    ws2.column_dimensions[get_column_letter(idx)].width=w
ws2.freeze_panes="A2"

out="/home/user/meet/italiamarket_catalog_full_2026.xlsx"
wb.save(out)
print("Saved",out,"| products:",len(rows))
