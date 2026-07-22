import pdfplumber, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

PDF = "/root/.claude/uploads/087eb6a4-9792-59b6-af15-064779936db4/4b7550f3-______________2026.pdf"

with pdfplumber.open(PDF) as pdf:
    t = pdf.pages[0].extract_text()
prod_lines = [l for l in t.split("\n") if re.search(r'\d+\.\d{2} EUR', l)]

# EN translations of the Bulgarian product names
EN = {
 "Сос Босилек": "Basil Sauce",
 "Сос Аматричана": "Amatriciana Sauce",
 "Сос ала Сицилиана": "Sauce alla Siciliana",
 "Сос каперси, маслини и аншоа": "Sauce with Capers, Olives and Anchovies",
 "Сос ала Норма": "Sauce alla Norma",
 "Сос Вентричина": "Ventricina Sauce",
 "Сос с Манатарки": "Sauce with Porcini Mushrooms",
 "Сос с Маслини": "Sauce with Olives",
 "Сос Арабиата": "Arrabbiata Sauce",
 "Сос Салсича": "Salsiccia (Sausage) Sauce",
 "Гъби микс": "Mushroom Mix",
 "Рулца от моркови и тиквички в маслиново масло": "Carrot and Zucchini Rolls in Olive Oil",
 "Люти чушки с пълнеж от мус от треска": "Hot Peppers stuffed with Cod Mousse",
 "Ръчно гриловани сърца от артишок с мента": "Hand-grilled Artichoke Hearts with Mint",
 "Сос качо е пепе": "Cacio e Pepe Sauce",
 "Песто Форте": "Pesto Forte",
 "Песто с маслини Лечино": "Pesto with Leccino Olives",
 "Песто с гъби и трюфели": "Pesto with Mushrooms and Truffles",
 "Песто с артишок": "Pesto with Artichoke",
 "Песто брускета": "Bruschetta Pesto",
 "Песто Пиканте": "Pesto Piccante (Spicy)",
 "Песто Класико с чесън": "Classic Pesto with Garlic",
 "Зелени маслини в саламура Сицилия": "Green Olives in Brine (Sicily)",
 "Коктейл маслини в саламура": "Cocktail Olives in Brine",
 "Сос Маринара": "Marinara Sauce",
 "Сос с маслини": "Sauce with Olives",
 "Сос с босилек": "Sauce with Basil",
 "Сос Традиционале": "Traditionale Sauce",
 'Ектра върджин маслиново масло "Opera Mastra"': 'Extra Virgin Olive Oil "Opera Mastra"',
 'Ектра върджин маслиново масло "TERRE DELL\'ABBAZIA"': 'Extra Virgin Olive Oil "TERRE DELL\'ABBAZIA"',
 'Екстра върджин маслиново масло " L\'O OLIO PREMIUM"': 'Extra Virgin Olive Oil "L\'O OLIO PREMIUM"',
 "Екстра върджин маслиново масло с пеперончино": "Extra Virgin Olive Oil with Peperoncino (Chili)",
 "Екстра върджин маслиново масло с босилек": "Extra Virgin Olive Oil with Basil",
 "Екстра върджин маслиново масло с бял трюфел": "Extra Virgin Olive Oil with White Truffle",
}
BRAND_EN = {"Урсини": "Ursini"}

# weight token: digits + unit (гр. / кг / мл / г / л)
row_re = re.compile(r'^(\S+)\s+(\d+\s*(?:гр\.|кг|мл|г\.?|л))\s+(.+?)\s+(Урсини)\s+бр\.\s+([\d.]+)\s+EUR\s+([\d.]+)\s+EUR$')

def norm_weight(w):
    w = w.replace(" ", "")
    w = w.replace("гр.", " g").replace("кг", " kg").replace("мл", " ml").replace("л", " l").replace("г", " g")
    return w.strip()

rows = []
unmatched = []
for l in prod_lines:
    m = row_re.match(l)
    if not m:
        unmatched.append(l); continue
    art, weight, name_bg, brand_bg, p_no, p_with = m.groups()
    name_bg = name_bg.strip()
    name_en = EN.get(name_bg, "??? "+name_bg)
    rows.append({
        "art": art,
        "name_en": name_en,
        "name_bg": name_bg,
        "weight": norm_weight(weight),
        "brand": BRAND_EN.get(brand_bg, brand_bg),
        "p_no": float(p_no),
        "p_with": float(p_with),
    })

print("Parsed:", len(rows), "Unmatched:", len(unmatched))
for u in unmatched: print("  UNMATCHED:", u)
missing = [r["name_en"] for r in rows if r["name_en"].startswith("???")]
print("Missing translations:", missing)
# VAT sanity
for r in rows:
    exp = round(r["p_no"]*1.2, 2)
    if abs(exp - r["p_with"]) > 0.02:
        print("  VAT mismatch:", r["art"], r["p_no"], r["p_with"], "exp", exp)

# ---- Build workbook ----
wb = Workbook()
ws = wb.active
ws.title = "Catalog"
headers = ["ART № (SKU)", "Product Name (EN)", "Product Name (BG)", "Weight", "Brand",
           "Price excl. VAT (EUR)", "Price incl. VAT (EUR)"]
ws.append(headers)

hdr_fill = PatternFill("solid", fgColor="1F4E78")
hdr_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
for c in ws[1]:
    c.fill = hdr_fill; c.font = hdr_font; c.border = border
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for i, r in enumerate(rows, start=2):
    ws.cell(i,1,r["art"])
    ws.cell(i,2,r["name_en"])
    ws.cell(i,3,r["name_bg"])
    ws.cell(i,4,r["weight"])
    ws.cell(i,5,r["brand"])
    ws.cell(i,6,r["p_no"])
    ws.cell(i,7,r["p_with"])
    for col in range(1,8):
        cell = ws.cell(i,col)
        cell.font = Font(name="Arial", size=10)
        cell.border = border
        cell.alignment = Alignment(vertical="center", wrap_text=(col in (2,3)))
        if col in (6,7):
            cell.number_format = '#,##0.00 "EUR"'
            cell.alignment = Alignment(horizontal="right", vertical="center")
    if i % 2 == 0:
        for col in range(1,8):
            ws.cell(i,col).fill = PatternFill("solid", fgColor="F2F6FB")

widths = [14, 42, 42, 10, 10, 20, 20]
for idx,w in enumerate(widths, start=1):
    ws.column_dimensions[chr(64+idx)].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:G{len(rows)+1}"

out = "/home/user/meet/italiamarket_catalog_2026.xlsx"
wb.save(out)
print("Saved:", out, "rows:", len(rows))
